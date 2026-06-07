#!/usr/bin/env python3

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def clean_value(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text


def first_non_empty_row(values: list[list[str]]) -> list[str]:
    for row in values:
        if any(cell for cell in row):
            return row
    return []


def audit_xlsx(source: Path) -> dict[str, Any]:
    workbook_formula = load_workbook(source, data_only=False)
    workbook_values = load_workbook(source, data_only=True)

    visible_sheets: list[str] = []
    hidden_sheets: list[str] = []
    sheet_reports: list[dict[str, Any]] = []
    warnings: list[str] = []

    for sheet_name in workbook_formula.sheetnames:
        ws_formula = workbook_formula[sheet_name]
        ws_values = workbook_values[sheet_name]

        if ws_formula.sheet_state == "hidden":
            hidden_sheets.append(sheet_name)
            warnings.append(f"Hidden sheet present: {sheet_name}")
        else:
            visible_sheets.append(sheet_name)

        merged_ranges = [str(rng) for rng in ws_formula.merged_cells.ranges]
        if merged_ranges:
            warnings.append(f"Merged cells present on sheet: {sheet_name}")

        formula_cells = 0
        displayed_value_cells = 0
        sampled_rows: list[list[str]] = []

        max_row = ws_formula.max_row or 0
        max_col = ws_formula.max_column or 0
        sample_limit = min(max_row, 12)
        column_limit = min(max_col, 12)

        for row_idx in range(1, sample_limit + 1):
            row_values: list[str] = []
            for col_idx in range(1, column_limit + 1):
                cell_formula = ws_formula.cell(row=row_idx, column=col_idx).value
                cell_value = ws_values.cell(row=row_idx, column=col_idx).value
                if isinstance(cell_formula, str) and cell_formula.startswith("="):
                    formula_cells += 1
                if cell_value not in (None, ""):
                    displayed_value_cells += 1
                row_values.append(clean_value(cell_value if cell_value is not None else cell_formula))
            sampled_rows.append(row_values)

        header_row = first_non_empty_row(sampled_rows)
        if not header_row:
            warnings.append(f"No representative header row found on sheet: {sheet_name}")

        if formula_cells > 0 and displayed_value_cells == 0:
            warnings.append(f"Formula-heavy sheet without displayed values in sample: {sheet_name}")

        sheet_reports.append(
            {
                "sheet_name": sheet_name,
                "sheet_state": ws_formula.sheet_state,
                "used_rows": max_row,
                "used_columns": max_col,
                "merged_cell_ranges": merged_ranges,
                "formula_cell_count_in_sample": formula_cells,
                "displayed_value_cell_count_in_sample": displayed_value_cells,
                "representative_header_row": header_row,
                "representative_sample_rows": [
                    row for row in sampled_rows[1:4] if any(cell for cell in row)
                ],
            }
        )

    if not sheet_reports:
        disposition = "fail"
        eligibility = "blocked"
        warnings.append("Workbook contained no readable sheets")
    elif warnings:
        disposition = "warning"
        eligibility = "eligible_with_review"
    else:
        disposition = "pass"
        eligibility = "eligible"

    return {
        "source_filename": source.name,
        "source_sha256": sha256(source),
        "format": "xlsx",
        "sheet_names": workbook_formula.sheetnames,
        "visible_sheets": visible_sheets,
        "hidden_sheets": hidden_sheets,
        "sheets": sheet_reports,
        "warnings": sorted(set(warnings)),
        "disposition": disposition,
        "normalisation_eligibility": eligibility,
        "uncertainty_notes": [
            "Workbook semantics beyond sampled structure still require human judgement when warnings are present."
        ],
    }


def slide_title(slide) -> str | None:
    title_shape = slide.shapes.title
    if title_shape and getattr(title_shape, "text", "").strip():
        return title_shape.text.strip()
    return None


def audit_pptx(source: Path) -> dict[str, Any]:
    presentation = Presentation(source)
    slide_reports: list[dict[str, Any]] = []
    warnings: list[str] = []

    for index, slide in enumerate(presentation.slides, start=1):
        text_blocks: list[str] = []
        image_count = 0
        object_count = len(slide.shapes)

        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if hasattr(shape, "has_text_frame") and shape.has_text_frame:
                text = "\n".join(
                    paragraph.text.strip()
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ).strip()
                if text:
                    text_blocks.append(text)

        notes_available = False
        notes_excerpt = ""
        try:
            notes_texts: list[str] = []
            for shape in slide.notes_slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        notes_texts.append(text)
            if notes_texts:
                notes_available = True
                notes_excerpt = " | ".join(notes_texts[:3])[:500]
        except Exception:
            notes_available = False

        layout_warning = False
        if image_count > 0 and len(text_blocks) <= 1:
            layout_warning = True
            warnings.append(f"Image-heavy or layout-dependent slide: {index}")
        if not text_blocks:
            layout_warning = True
            warnings.append(f"Slide has no extracted text blocks: {index}")

        slide_reports.append(
            {
                "slide_number": index,
                "title": slide_title(slide),
                "text_blocks": text_blocks,
                "text_block_count": len(text_blocks),
                "object_count": object_count,
                "image_count": image_count,
                "notes_available": notes_available,
                "notes_excerpt": notes_excerpt,
                "layout_dependent_warning": layout_warning,
            }
        )

    if len(presentation.slides) == 0:
        disposition = "fail"
        eligibility = "blocked"
        warnings.append("Presentation contains zero slides")
    elif warnings:
        disposition = "warning"
        eligibility = "eligible_with_review"
    else:
        disposition = "pass"
        eligibility = "eligible"

    return {
        "source_filename": source.name,
        "source_sha256": sha256(source),
        "format": "pptx",
        "slide_count": len(presentation.slides),
        "slides": slide_reports,
        "warnings": sorted(set(warnings)),
        "disposition": disposition,
        "normalisation_eligibility": eligibility,
        "uncertainty_notes": [
            "Visual layout meaning still requires human judgement when slide warnings are present."
        ],
    }


def normalise_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def audit_pdf(source: Path, docling_json_path: Path) -> dict[str, Any]:
    data = json.loads(docling_json_path.read_text(encoding="utf-8"))
    pages = data.get("pages", {})
    texts = data.get("texts", [])
    tables = data.get("tables", [])
    pictures = data.get("pictures", [])
    body_children = data.get("body", {}).get("children", [])

    warnings: list[str] = []
    page_reports: list[dict[str, Any]] = []

    page_numbers: list[int] = []
    if isinstance(pages, dict):
        for key in sorted(pages, key=lambda x: int(str(x)) if str(x).isdigit() else str(x)):
            page = pages[key]
            page_no = int(page.get("page_no", key))
            page_numbers.append(page_no)
    elif isinstance(pages, list):
        for idx, page in enumerate(pages, start=1):
            page_no = int(page.get("page_no", idx))
            page_numbers.append(page_no)

    text_refs_by_page: defaultdict[int, list[dict[str, Any]]] = defaultdict(list)
    bounding_box_presence = 0
    provenance_presence = 0
    normalized_counter: Counter[str] = Counter()

    for item in texts:
        prov = item.get("prov") or []
        if prov:
            provenance_presence += 1
        for prov_entry in prov:
            page_no = int(prov_entry.get("page_no", 0))
            if page_no:
                text_refs_by_page[page_no].append(
                    {
                        "label": item.get("label"),
                        "text": item.get("text", ""),
                        "bbox": prov_entry.get("bbox"),
                    }
                )
                if prov_entry.get("bbox"):
                    bounding_box_presence += 1

        normalized = normalise_text(item.get("text", ""))
        if normalized:
            normalized_counter[normalized] += 1

    duplicate_blocks = [text for text, count in normalized_counter.items() if count > 1 and len(text) > 20]
    if duplicate_blocks:
        warnings.append(f"Duplicate text blocks detected: {len(duplicate_blocks)}")

    for page_no in page_numbers:
        entries = text_refs_by_page.get(page_no, [])
        if not entries:
            warnings.append(f"Suspicious empty page or missing text provenance: page {page_no}")
        left_positions = sorted(
            {
                round(entry["bbox"].get("l", 0), 2)
                for entry in entries
                if entry.get("bbox")
            }
        )
        layout_risk = len(left_positions) >= 3 and (max(left_positions) - min(left_positions) > 120)
        if layout_risk:
            warnings.append(f"Potential layout-risk page detected: page {page_no}")

        page_reports.append(
            {
                "page_number": page_no,
                "text_block_count": len(entries),
                "left_position_samples": left_positions[:10],
                "layout_risk_warning": layout_risk,
            }
        )

    ordered_document_items = [
        child.get("$ref")
        for child in body_children
        if isinstance(child, dict) and child.get("$ref")
    ]

    if not page_numbers:
        warnings.append("No page count exposed in Docling JSON")
        disposition = "fail"
        eligibility = "blocked"
    elif provenance_presence == 0 or not ordered_document_items:
        warnings.append("Reading-order evidence incomplete in Docling JSON")
        disposition = "fail"
        eligibility = "blocked"
    elif warnings:
        disposition = "warning"
        eligibility = "eligible_with_review"
    else:
        disposition = "pass"
        eligibility = "eligible"

    return {
        "source_filename": source.name,
        "source_sha256": sha256(source),
        "format": "pdf",
        "docling_json_path": str(docling_json_path),
        "page_count": len(page_numbers),
        "page_numbers": page_numbers,
        "ordered_document_item_sequence": ordered_document_items,
        "text_block_count": len(texts),
        "provenance_present_for_text_items": provenance_presence,
        "bounding_box_provenance_count": bounding_box_presence,
        "table_count": len(tables),
        "picture_count": len(pictures),
        "pages": page_reports,
        "warnings": sorted(set(warnings)),
        "disposition": disposition,
        "normalisation_eligibility": eligibility,
        "uncertainty_notes": [
            "Docling JSON supports a page-aware structural audit, not automatic semantic certainty.",
            "Multi-column or layout-heavy PDFs may still require bounded human review when warnings are present.",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit extraction quality for XLSX, PPTX or PDF inputs."
    )
    parser.add_argument("--source", required=True, help="Path to the source file.")
    parser.add_argument(
        "--format",
        required=True,
        choices=["xlsx", "pptx", "pdf"],
        help="Format-specific audit module to run.",
    )
    parser.add_argument("--output-json", required=True, help="Path for audit JSON output.")
    parser.add_argument(
        "--docling-json",
        help="Required for PDF mode: path to Docling structured JSON.",
    )

    args = parser.parse_args()

    source = Path(args.source).resolve()
    output_json = Path(args.output_json).resolve()

    if not source.is_file():
        raise FileNotFoundError(f"Source file not found: {source}")

    if args.format == "xlsx":
        result = audit_xlsx(source)
    elif args.format == "pptx":
        result = audit_pptx(source)
    else:
        if not args.docling_json:
            raise ValueError("--docling-json is required for pdf mode")
        docling_json = Path(args.docling_json).resolve()
        if not docling_json.is_file():
            raise FileNotFoundError(f"Docling JSON not found: {docling_json}")
        result = audit_pdf(source, docling_json)

    write_json(output_json, result)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
