#!/usr/bin/env python3

from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
import time
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore

try:
    from pptx import Presentation
except Exception:  # pragma: no cover
    Presentation = None  # type: ignore


LAB_ROOT = Path(__file__).resolve().parents[2]
PROJECT_ID = "01-doc-to-spec-pilot"
PROJECT_DIR = LAB_ROOT / "projects" / PROJECT_ID
DOC_TO_SPEC_VENV_PYTHON = PROJECT_DIR / ".venv" / "bin" / "python"
DEFAULT_FIXTURE_DIR = PROJECT_DIR / "fixtures" / "circuitfit-synthetic-enterprise-evidence-pack-v0.1"
STEP16_LOG_DIR = LAB_ROOT / "logs" / "step16"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def safe_slug(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in value)


def format_classification(path: Path) -> str:
    ext = path.suffix.lower()
    mapping = {
        ".md": "text",
        ".txt": "text",
        ".csv": "csv",
        ".eml": "email",
        ".docx": "docx",
        ".xlsx": "xlsx",
        ".pptx": "pptx",
        ".pdf": "pdf",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image",
    }
    return mapping.get(ext, "unsupported")


def estimate_file_work_units(path: Path) -> dict[str, Any]:
    ext = path.suffix.lower()
    estimate: dict[str, Any] = {
        "file": path.name,
        "format": ext,
        "files": 1,
        "pages": 0,
        "slides": 0,
        "sheets": 0,
        "images": 0,
        "ocr_needed": False,
        "weighted_units": 1,
    }

    if ext == ".xlsx" and load_workbook is not None:
        workbook = load_workbook(path, read_only=True, data_only=True)
        estimate["sheets"] = len(workbook.sheetnames)
        estimate["weighted_units"] = max(2, len(workbook.sheetnames) * 2)
    elif ext == ".pptx" and Presentation is not None:
        deck = Presentation(path)
        estimate["slides"] = len(deck.slides)
        image_count = 0
        for slide in deck.slides:
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) == 13:
                    image_count += 1
        estimate["images"] = image_count
        estimate["weighted_units"] = max(2, len(deck.slides) * 2 + image_count)
    elif ext == ".pdf":
        estimate["ocr_needed"] = True
        estimate["weighted_units"] = 6
    elif ext in {".png", ".jpg", ".jpeg"}:
        estimate["images"] = 1
        estimate["ocr_needed"] = True
        estimate["weighted_units"] = 4
    elif ext == ".docx":
        estimate["weighted_units"] = 2
    elif ext == ".eml":
        estimate["weighted_units"] = 1
    return estimate


def stage_record(name: str, status: str, started: float, **extra: Any) -> dict[str, Any]:
    record = {
        "stage": name,
        "status": status,
        "completed_utc": utc_now(),
        "duration_seconds": round(time.perf_counter() - started, 3),
    }
    record.update(extra)
    return record


def resolve_project(args: argparse.Namespace) -> str:
    project = args.project_option or args.project
    if project not in {"circuitfit"}:
        raise ValueError("A supported project must be supplied via positional argument or --project.")
    return project


def collect_files(input_dir: Path) -> list[Path]:
    return sorted(path for path in input_dir.iterdir() if path.is_file() and path.name != ".gitkeep")


def summarise_files(files: list[Path]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    work_units = [estimate_file_work_units(path) for path in files]
    files_summary = []
    for path, estimate in zip(files, work_units):
        files_summary.append(
            {
                "filename": path.name,
                "extension": path.suffix.lower(),
                "classification": format_classification(path),
                "source_sha256": sha256(path),
                "size_bytes": path.stat().st_size,
                "work_unit_estimate": estimate,
            }
        )
    return files_summary, work_units


def validate_manifest(manifest: dict[str, Any], expected_project: str) -> tuple[bool, list[str]]:
    warnings: list[str] = []
    required_keys = [
        "run_id",
        "project_slug",
        "input_directory",
        "files",
        "pipeline_state",
        "resume_stage",
        "human_gate",
    ]
    for key in required_keys:
        if key not in manifest:
            warnings.append(f"Missing required manifest key: {key}")
    if manifest.get("project_slug") != expected_project:
        warnings.append("Manifest project does not match requested project.")
    if not isinstance(manifest.get("files"), list) or not manifest.get("files"):
        warnings.append("Manifest file inventory is missing or empty.")
    human_gate = manifest.get("human_gate", {})
    if not isinstance(human_gate, dict) or "approval_phrase" not in human_gate:
        warnings.append("Manifest human gate is missing approval phrase.")
    return not warnings, warnings


def source_hash_validation(manifest: dict[str, Any], input_dir: Path) -> dict[str, Any]:
    mismatches: list[dict[str, str]] = []
    files_in_manifest = manifest.get("files", [])
    manifest_by_name = {item["filename"]: item for item in files_in_manifest if "filename" in item}
    current_files = collect_files(input_dir)
    current_names = {path.name for path in current_files}
    manifest_names = set(manifest_by_name)

    missing_from_disk = sorted(manifest_names - current_names)
    new_on_disk = sorted(current_names - manifest_names)

    for path in current_files:
        manifest_item = manifest_by_name.get(path.name)
        if not manifest_item:
            continue
        current_hash = sha256(path)
        if current_hash != manifest_item.get("source_sha256"):
            mismatches.append(
                {
                    "filename": path.name,
                    "manifest_sha256": str(manifest_item.get("source_sha256")),
                    "current_sha256": current_hash,
                }
            )

    unchanged = not mismatches and not missing_from_disk and not new_on_disk
    return {
        "unchanged": unchanged,
        "mismatches": mismatches,
        "missing_from_disk": missing_from_disk,
        "new_on_disk": new_on_disk,
    }


def find_latest_valid_manifest(project: str, input_dir: Path) -> tuple[Path | None, dict[str, Any] | None, list[str]]:
    warnings: list[str] = []
    if not STEP16_LOG_DIR.is_dir():
        return None, None, [f"Manifest search directory not found: {STEP16_LOG_DIR}"]
    candidates = sorted(
        STEP16_LOG_DIR.glob(f"pmcp-{project}-*/pipeline_manifest.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for candidate in candidates:
        try:
            manifest = read_json(candidate)
        except json.JSONDecodeError as exc:
            warnings.append(f"Invalid JSON manifest skipped: {candidate} ({exc})")
            continue
        valid, manifest_warnings = validate_manifest(manifest, project)
        if not valid:
            warnings.extend(f"{candidate}: {warning}" for warning in manifest_warnings)
            continue
        if Path(manifest["input_directory"]).resolve() != input_dir:
            warnings.append(f"Skipped manifest with different input directory: {candidate}")
            continue
        return candidate, manifest, warnings
    return None, None, warnings


def approval_recorded(manifest: dict[str, Any], provided_approval_phrase: str | None) -> bool:
    human_gate = manifest.get("human_gate", {})
    required_phrase = human_gate.get("approval_phrase")
    recorded = manifest.get("approval_record", {})
    recorded_phrase = recorded.get("approval_phrase")
    recorded_status = recorded.get("status")
    if recorded_status == "approved" and recorded_phrase == required_phrase:
        return True
    return bool(required_phrase and provided_approval_phrase == required_phrase)


def determine_route(
    manifest: dict[str, Any] | None,
    manifest_path: Path | None,
    hash_validation: dict[str, Any] | None,
    provided_approval_phrase: str | None,
    rerun_requested: bool,
) -> dict[str, Any]:
    if manifest is None or manifest_path is None:
        return {
            "route_decision": "start_new_run",
            "next_stage": "preflight",
            "plain_english": "No valid prior manifest was found, so the workflow would start at preflight.",
            "should_rerun_ingestion": True,
            "openspec_blocked": True,
        }

    if rerun_requested:
        return {
            "route_decision": "rerun_requested",
            "next_stage": "preflight",
            "plain_english": "The user explicitly requested a rerun, so the workflow would restart at preflight.",
            "should_rerun_ingestion": True,
            "openspec_blocked": True,
        }

    if hash_validation and not hash_validation["unchanged"]:
        return {
            "route_decision": "source_hash_changed",
            "next_stage": "preflight",
            "plain_english": "Upstream source hashes changed, so completed ingestion stages would not be reused.",
            "should_rerun_ingestion": True,
            "openspec_blocked": True,
        }

    state = manifest.get("pipeline_state")
    resume_stage = manifest.get("resume_stage")
    approval_found = approval_recorded(manifest, provided_approval_phrase)

    if state == "awaiting_human_approval_before_normalisation":
        if approval_found:
            return {
                "route_decision": "resume_after_evidence_approval",
                "next_stage": resume_stage or "normalisation",
                "plain_english": "The evidence-eligibility approval is recorded, so the next invocation would resume at normalisation without rerunning ingestion.",
                "should_rerun_ingestion": False,
                "openspec_blocked": True,
            }
        return {
            "route_decision": "show_evidence_gate",
            "next_stage": "human_gate",
            "plain_english": "The workflow is waiting at the evidence-eligibility gate and still needs the recorded approval phrase before normalisation may begin.",
            "should_rerun_ingestion": False,
            "openspec_blocked": True,
        }

    if state == "awaiting_human_review_of_normalised_evidence_and_conflicts":
        return {
            "route_decision": "show_conflict_review_gate",
            "next_stage": "conflict_review_gate",
            "plain_english": "The workflow is waiting for human review of normalised evidence and conflicts.",
            "should_rerun_ingestion": False,
            "openspec_blocked": True,
        }

    if state == "approved_for_openspec_proposal":
        return {
            "route_decision": "resume_openspec_proposal",
            "next_stage": "openspec_proposal",
            "plain_english": "Conflict review was already approved, so the next resumable phase would be the OpenSpec proposal stage.",
            "should_rerun_ingestion": False,
            "openspec_blocked": False,
        }

    return {
        "route_decision": "state_not_yet_modelled",
        "next_stage": resume_stage,
        "plain_english": "A prior manifest was found, but its pipeline state does not match a supported resumable route yet.",
        "should_rerun_ingestion": False,
        "openspec_blocked": True,
    }


def run_step16_pipeline(project: str, input_dir: Path, output_dir: Path | None) -> dict[str, Any]:
    if not DOC_TO_SPEC_VENV_PYTHON.is_file():
        raise FileNotFoundError(f"Docling environment python not found: {DOC_TO_SPEC_VENV_PYTHON}")

    run_id = f"pmcp-{project}-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
    resolved_output_dir = output_dir.resolve() if output_dir else LAB_ROOT / "logs" / "step16" / safe_slug(run_id)
    resolved_output_dir.mkdir(parents=True, exist_ok=True)
    extraction_dir = resolved_output_dir / "extracted"
    extraction_report_json = resolved_output_dir / "extraction_report.json"
    audit_dir = resolved_output_dir / "audit_results"
    manifest_path = resolved_output_dir / "pipeline_manifest.json"

    manifest: dict[str, Any] = {
        "run_id": run_id,
        "project_id": PROJECT_ID,
        "project_slug": project,
        "triggered_at_utc": utc_now(),
        "git_commit": subprocess.check_output(["git", "rev-parse", "--short", "HEAD"], cwd=LAB_ROOT, text=True).strip(),
        "agent_name": "product-manager-copilot",
        "skill_name": "product-manager-copilot-doc-to-spec",
        "model_identifier": "openai/gpt-5.4",
        "input_directory": str(input_dir),
        "local_output_directory": str(resolved_output_dir),
        "stages": [],
        "files": [],
        "warnings": [],
        "blocked_evidence": [],
        "deferred_formats": [".msg"],
        "reviewer_attention_items": [],
        "pipeline_state": "running",
        "resume_stage": None,
    }

    files = collect_files(input_dir)

    preflight_started = time.perf_counter()
    files_summary, work_units = summarise_files(files)
    manifest["files"] = files_summary
    manifest["format_counts"] = dict(Counter(item["classification"] for item in files_summary))
    manifest["work_unit_estimate"] = {
        key: sum(item[key] for item in work_units)
        for key in ["files", "pages", "slides", "sheets", "images", "weighted_units"]
    }
    manifest["work_unit_estimate"]["ocr_needed_files"] = sum(1 for item in work_units if item["ocr_needed"])
    manifest["stages"].append(stage_record("preflight", "completed", preflight_started))

    extraction_started = time.perf_counter()
    subprocess.run(
        [
            str(DOC_TO_SPEC_VENV_PYTHON),
            str(LAB_ROOT / "shared" / "scripts" / "extract-multiformat-evidence-pack.py"),
            "--input-dir",
            str(input_dir),
            "--output-dir",
            str(extraction_dir),
            "--report-json",
            str(extraction_report_json),
        ],
        cwd=LAB_ROOT,
        check=True,
    )
    extraction_report = read_json(extraction_report_json)
    manifest["extraction_report_json"] = str(extraction_report_json)
    manifest["stages"].append(
        stage_record(
            "extraction",
            "completed",
            extraction_started,
            files_completed=extraction_report.get("completed_count"),
            files_failed=extraction_report.get("failed_count"),
        )
    )

    extraction_results_by_file = {item["source_filename"]: item for item in extraction_report.get("results", [])}

    audit_started = time.perf_counter()
    audit_results: dict[str, Any] = {}
    for filename, fmt in [
        ("04_workout_requirements_matrix.xlsx", "xlsx"),
        ("05_mobile_wireframe_review.pptx", "pptx"),
        ("06_beta_user_feedback_summary.pdf", "pdf"),
    ]:
        source = input_dir / filename
        audit_output = audit_dir / f"{safe_slug(source.stem)}.{fmt}.audit.json"
        command = [
            sys.executable,
            str(LAB_ROOT / "shared" / "scripts" / "audit-multiformat-extraction-quality.py"),
            "--source",
            str(source),
            "--format",
            fmt,
            "--output-json",
            str(audit_output),
        ]
        if fmt == "pdf":
            docling_json = extraction_results_by_file[filename]["output_json"]
            command.extend(["--docling-json", str(docling_json)])
        subprocess.run(command, cwd=LAB_ROOT, check=True)
        audit_results[filename] = read_json(audit_output)

    manifest["audit_results"] = audit_results
    manifest["stages"].append(stage_record("quality_audit", "completed", audit_started))

    image_started = time.perf_counter()
    image_decisions = {
        "10_mobile_webapp_screenshot_feedback.png": {
            "disposition": "bounded_human_review",
            "normalisation_eligibility": "eligible_with_review",
            "human_review_requirement": "Confirm the exact screenshot-derived labels or statements used in the first normalisation pass.",
        },
        "09_whiteboard_workflow_photo.png": {
            "disposition": "blocked",
            "normalisation_eligibility": "blocked",
            "human_review_requirement": "Do not use whiteboard-derived assertions in the first OpenSpec experiment.",
        },
        ".msg": {
            "disposition": "deferred",
            "normalisation_eligibility": "deferred",
            "human_review_requirement": "Maintain explicit deferred route only; do not ingest silently.",
        },
    }
    manifest["image_and_deferred_decisions"] = image_decisions
    manifest["blocked_evidence"].append("09_whiteboard_workflow_photo.png")
    manifest["reviewer_attention_items"].append(
        "10_mobile_webapp_screenshot_feedback.png requires bounded human review before any screenshot-derived assertions may be normalised."
    )
    manifest["reviewer_attention_items"].extend(
        warning
        for result in audit_results.values()
        for warning in result.get("warnings", [])
    )
    manifest["stages"].append(stage_record("image_decisions", "completed", image_started))

    summary_started = time.perf_counter()
    eligibility_summary = []
    baseline_eligible = {
        "00_README_FIRST.md": "eligible",
        "01_email_product_owner_priority_change.eml": "eligible",
        "02_email_engineering_storage_privacy.eml": "eligible",
        "03_founder_review_meeting_minutes.docx": "eligible",
        "07_architecture_chat_export.txt": "eligible",
        "08_stale_feature_request_v0_8.docx": "eligible",
        "11_EXPECTED_CONFLICTS.md": "eligible",
        "12_EXPECTED_EXTRACTION_FEATURES.md": "eligible",
        "13_manifest.csv": "eligible",
    }
    for item in files_summary:
        filename = item["filename"]
        if filename in audit_results:
            eligibility = audit_results[filename]["normalisation_eligibility"]
            disposition = audit_results[filename]["disposition"]
            warnings = audit_results[filename].get("warnings", [])
        elif filename == "10_mobile_webapp_screenshot_feedback.png":
            eligibility = "eligible_with_review"
            disposition = "bounded_human_review"
            warnings = [image_decisions[filename]["human_review_requirement"]]
        elif filename == "09_whiteboard_workflow_photo.png":
            eligibility = "blocked"
            disposition = "blocked"
            warnings = [image_decisions[filename]["human_review_requirement"]]
        else:
            eligibility = baseline_eligible.get(filename, "blocked")
            disposition = "pass" if eligibility == "eligible" else "warning"
            warnings = []
        eligibility_summary.append(
            {
                "filename": filename,
                "classification": item["classification"],
                "disposition": disposition,
                "normalisation_eligibility": eligibility,
                "warnings": warnings,
            }
        )

    manifest["normalisation_eligibility_summary"] = eligibility_summary
    manifest["pipeline_state"] = "awaiting_human_approval_before_normalisation"
    manifest["resume_stage"] = "normalisation"
    manifest["stages"].append(stage_record("eligibility_summary", "completed", summary_started))

    human_gate_started = time.perf_counter()
    manifest["human_gate"] = {
        "approval_phrase": "APPROVE ELIGIBLE CIRCUITFIT EVIDENCE FOR NORMALISATION",
        "what_this_authorises": [
            "Normalisation of evidence explicitly marked eligible or eligible_with_review after the required human review.",
            "Progression to the next bounded Step 17 normalisation work.",
        ],
        "what_remains_blocked": [
            "09_whiteboard_workflow_photo.png",
            "Outlook .msg",
            "Any evidence file still marked blocked",
        ],
        "what_this_does_not_authorise": [
            "OpenSpec generation",
            "Application coding",
            "Test creation",
            "Deployment",
        ],
    }
    manifest["stages"].append(stage_record("human_gate", "completed", human_gate_started))

    write_json(manifest_path, manifest)
    return {
        "run_id": run_id,
        "manifest_path": str(manifest_path),
        "pipeline_state": manifest["pipeline_state"],
        "resume_stage": manifest["resume_stage"],
        "approval_phrase": manifest["human_gate"]["approval_phrase"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Run or resume the Product Manager CoPilot document-to-spec workflow."
    )
    parser.add_argument("project", nargs="?", help="Synthetic pilot project to process.")
    parser.add_argument("--project", dest="project_option", choices=["circuitfit"], help="Synthetic pilot project to process.")
    parser.add_argument(
        "--input-dir",
        default=str(DEFAULT_FIXTURE_DIR),
        help="Evidence-pack directory to process.",
    )
    parser.add_argument(
        "--output-dir",
        help="Optional local-only output directory. Defaults to logs/step16/<run-id>/ for new Step 16 runs.",
    )
    parser.add_argument(
        "--manifest",
        help="Optional explicit manifest path to inspect for resumable routing.",
    )
    parser.add_argument(
        "--resume",
        action="store_true",
        help="Inspect the latest valid manifest and route to the next resumable stage without rerunning completed stages unless required.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Show the resumable routing decision without executing the next stage.",
    )
    parser.add_argument(
        "--approval-phrase",
        help="Recorded human approval phrase for the next bounded phase when applicable.",
    )
    parser.add_argument(
        "--rerun",
        action="store_true",
        help="Explicitly rerun from preflight instead of reusing the latest valid manifest.",
    )
    args = parser.parse_args()

    project = resolve_project(args)
    input_dir = Path(args.input_dir).resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"Input directory not found: {input_dir}")

    if args.resume:
        manifest_warnings: list[str] = []
        manifest_path: Path | None = None
        manifest: dict[str, Any] | None = None
        if args.manifest:
            manifest_path = Path(args.manifest).resolve()
            manifest = read_json(manifest_path)
            valid, explicit_manifest_warnings = validate_manifest(manifest, project)
            manifest_warnings.extend(explicit_manifest_warnings)
            if not valid:
                manifest = None
        else:
            manifest_path, manifest, manifest_warnings = find_latest_valid_manifest(project, input_dir)

        hash_validation = source_hash_validation(manifest, input_dir) if manifest else None
        route = determine_route(manifest, manifest_path, hash_validation, args.approval_phrase, args.rerun)

        output = {
            "project": project,
            "resume_requested": True,
            "dry_run": bool(args.dry_run),
            "latest_valid_manifest_found": bool(manifest_path and manifest),
            "manifest_path": str(manifest_path) if manifest_path else None,
            "manifest_warnings": manifest_warnings,
            "approval_phrase_found": approval_recorded(manifest, args.approval_phrase) if manifest else False,
            "required_approval_phrase": manifest.get("human_gate", {}).get("approval_phrase") if manifest else None,
            "provided_approval_phrase": args.approval_phrase,
            "source_hash_validation": hash_validation,
            "completed_ingestion_stages_rerun": False,
            "next_routed_stage": route["next_stage"],
            "route_decision": route["route_decision"],
            "plain_english": route["plain_english"],
            "openspec_blocked": route["openspec_blocked"],
            "normalisation_performed": False,
            "application_code_written": False,
        }
        print(json.dumps(output, indent=2, ensure_ascii=False))
        return 0

    result = run_step16_pipeline(
        project=project,
        input_dir=input_dir,
        output_dir=Path(args.output_dir) if args.output_dir else None,
    )
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
